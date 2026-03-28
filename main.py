from fastapi import FastAPI, UploadFile, File
from fastapi.middleware.cors import CORSMiddleware
from dotenv import load_dotenv
import fitz  # PyMuPDF
from google import genai
import os
from pydantic import BaseModel
from pptx import Presentation
from fastapi.responses import FileResponse
import os


load_dotenv()
client = genai.Client(api_key=os.getenv("GEMINI_API_KEY"))

app = FastAPI()

app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:5173"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


# ==========================================
# 1. SINGLE LESSON PLAN (UPGRADED FOR PPTX)
# ==========================================
@app.post("/generate-lesson/{lesson_id}")
async def generate_lesson(lesson_id: int, file: UploadFile = File(...)):
    try:
        pdf_bytes = await file.read()
        doc = fitz.open(stream=pdf_bytes, filetype="pdf")
        extracted_text = "".join(page.get_text() for page in doc)

        prompt = f"""
        You are an expert Educational Curriculum Designer and a TED Talk Presentation Writer. 
        Create a highly detailed and engaging lesson plan for Week/Lesson #{lesson_id}.

        CRITICAL INSTRUCTIONS FOR CONTENT:
        1. DEEP DEFINITIONS: For EVERY topic, provide a clear, in-depth explanation.
        2. REAL-WORLD EXAMPLES: For EVERY topic, provide a relatable, real-life practical example to make it easy to understand.
        3. VISUAL GRAPHICS: Use relevant Emojis (e.g., 🌍📈💻⚙️) generously as visual icons to make the text engaging for presentations.

        Use strict, professional English. Format cleanly using Markdown.
        Include these sections:
        ## 🎯 Lesson Objectives
        ## 📚 Key Topics (Include Definitions & Real-World Examples here)
        ## 💡 Interactive Activities
        ## 📝 Homework / Assignment
        ## ⏱️ Quick End-of-Class Quiz

        Syllabus Text: {extracted_text}
        """
        response = client.models.generate_content(model='gemini-2.5-flash', contents=prompt)
        return {"content": response.text}
    except Exception as e:
        return {"error": str(e)}

# ==========================================
# 2. FULL COURSE ROADMAP
# ==========================================
@app.post("/generate-course-plan/{total_weeks}")
async def generate_course_plan(total_weeks: int, file: UploadFile = File(...)):
    try:
        pdf_bytes = await file.read()
        doc = fitz.open(stream=pdf_bytes, filetype="pdf")
        extracted_text = "".join(page.get_text() for page in doc)

        prompt = f"""
        You are an expert Academic Planner. Divide and map the ENTIRE syllabus across exactly {total_weeks} weeks evenly.
        Use strict, professional English.
        For EACH week, use this structure:
        ## Week [Number]: [Main Topic]
        - **Topics to Cover:** ...
        - **Suggested Activity:** ...

        Syllabus Text: {extracted_text}
        """
        response = client.models.generate_content(model='gemini-2.5-flash', contents=prompt)
        return {"content": response.text}
    except Exception as e:
        return {"error": str(e)}


# ==========================================
# 3. NAYA: SMART EXAM GENERATOR
# ==========================================
@app.post("/generate-exam/{difficulty}")
async def generate_exam(difficulty: str, file: UploadFile = File(...)):
    try:
        pdf_bytes = await file.read()
        doc = fitz.open(stream=pdf_bytes, filetype="pdf")
        extracted_text = "".join(page.get_text() for page in doc)

        prompt = f"""
                You are an expert University Professor. Create a comprehensive exam based on the entire syllabus provided.
                The overall difficulty level MUST be: {difficulty}.

                CRITICAL FORMATTING RULES:
                - You MUST use double spaces/newlines between every single question so they don't merge into one paragraph.
                - For Section A (MCQs), format EXACTLY like this using bullet points:
                  **1. [Question Text]**
                  * A) [Option 1]
                  * B) [Option 2]
                  * C) [Option 3]
                  * D) [Option 4]
                - For Section B and C, leave a blank line after the question text.

                Include these sections:
                ## 📝 Term Examination
                ### Section A: Multiple Choice (5 Questions)
                ### Section B: Short Answer Questions (3 Questions)
                ### Section C: Analytical/Essay Question (1 Complex Question)

                ---
                ## 🔐 Teacher's Answer Key
                (Provide correct answers and a brief grading guide here).

                Syllabus Text: {extracted_text}
                """
        response = client.models.generate_content(model='gemini-2.5-flash', contents=prompt)
        return {"content": response.text}
    except Exception as e:
        return {"error": str(e)}







# ==========================================
# 4. NAYA: ASSIGNMENT & RUBRIC BUILDER
# ==========================================
@app.post("/generate-assignment/{week_id}")
async def generate_assignment(week_id: int, file: UploadFile = File(...)):
    try:
        pdf_bytes = await file.read()
        doc = fitz.open(stream=pdf_bytes, filetype="pdf")
        extracted_text = "".join(page.get_text() for page in doc)

        prompt = f"""
                You are an expert University Professor. Create a practical assignment based strictly on the topics of Week/Lesson #{week_id} from the syllabus.

                CRITICAL FORMATTING RULES:
                - Add a blank line between every paragraph and heading.
                - You MUST generate a proper Markdown Table for the grading rubric.
                - Add a blank empty line BEFORE and AFTER the table to ensure it renders correctly.

                Include these exact sections:
                ## 📋 Assignment Brief (Week {week_id})
                (Clearly define the task, scenario, and deliverables for the students).

                ## 🎯 Grading Rubric
                | Criteria | Excellent (100%) | Average (50%) | Poor (0%) |
                | :--- | :--- | :--- | :--- |
                | **Logic & Accuracy** | (Describe) | (Describe) | (Describe) |
                | **Formatting & Quality** | (Describe) | (Describe) | (Describe) |
                | **Completeness** | (Describe) | (Describe) | (Describe) |

                Syllabus Text: {extracted_text}
                """
        response = client.models.generate_content(model='gemini-2.5-flash', contents=prompt)
        return {"content": response.text}
    except Exception as e:
        return {"error": str(e)}


# ==========================================
# 5. NAYA: 1-CLICK PPTX GENERATOR
# ==========================================
class PPTXRequest(BaseModel):
    text: str
    title: str

@app.post("/generate-pptx")
async def generate_pptx(request: PPTXRequest):
    try:
        prs = Presentation()

        # 1. Title Slide (Pehli Slide)
        title_slide_layout = prs.slide_layouts[0]
        slide = prs.slides.add_slide(title_slide_layout)
        title = slide.shapes.title
        subtitle = slide.placeholders[1]
        title.text = request.title
        subtitle.text = "Generated by LearnSync AI"

        # 2. Markdown Text ko Sections (Headings) mein torna
        # Hum '## ' se split kar rahe hain kyunke hamari headings H2 hain
        sections = request.text.split("## ")

        for section in sections[1:]:  # Pehla hissa khali ho sakta hai isliye [1:]
            lines = section.strip().split('\n')
            if not lines: continue

            # Pehli line slide ka title hogi, baqi sab body
            slide_title = lines[0].strip()
            slide_body = "\n".join(lines[1:]).strip()

            # Markdown ke extra stars (**) remove karna clean look ke liye
            slide_body = slide_body.replace("**", "")

            # Content Slide Add Karna
            bullet_slide_layout = prs.slide_layouts[1]
            slide = prs.slides.add_slide(bullet_slide_layout)

            slide.shapes.title.text = slide_title
            slide.shapes.placeholders[1].text = slide_body

        # File Save Karna aur Return Karna
        file_path = "LearnSync_Presentation.pptx"
        prs.save(file_path)

        return FileResponse(
            path=file_path,
            filename=f"{request.title}.pptx",
            media_type='application/vnd.openxmlformats-officedocument.presentationml.presentation'
        )

    except Exception as e:
        return {"error": str(e)}


# ==========================================
# 6. STUDENT PORTAL: MASTER GENERATOR
# ==========================================
@app.post("/student/generate/{mode}")
async def generate_student_content(mode: str, file: UploadFile = File(...)):
    try:
        pdf_bytes = await file.read()
        doc = fitz.open(stream=pdf_bytes, filetype="pdf")
        extracted_text = "".join(page.get_text() for page in doc)

        if mode == "plan":
            prompt = f"""
            You are an expert tutor. Create a comprehensive 'Standard Study Plan' from the syllabus.
            For EVERY topic, provide:
            1. **Detailed Explanation:** Explain the topic thoroughly so the student understands the core concept deeply.
            2. **Real-Life Example:** A highly relatable, everyday example.
            3. **YouTube Link:** Use this exact format: [📺 Watch Tutorial](https://www.youtube.com/results?search_query=[Replace+With+Topic+Name+With+Plus+Signs])

            Syllabus: {extracted_text}
            """

        elif mode == "panic":
            prompt = f"""
            🚨 PANIC MODE ACTIVATED! The student has an exam tomorrow. Apply the 80/20 rule. 
            Extract ONLY the top 20% absolute most critical, high-yield topics from the syllabus.
            For EACH of these critical topics, provide:
            1. **Detailed Explanation:** Explain the topic thoroughly but directly.
            2. **Real-Life Example:** A highly relatable, practical example to lock it in memory.
            3. **YouTube Link:** Use this exact format: [📺 Watch Tutorial](https://www.youtube.com/results?search_query=[Replace+With+Topic+Name+With+Plus+Signs])

            Syllabus: {extracted_text}
            """

        elif mode == "crash":
            prompt = f"""
            You are creating a quick Crash Course. 
            For the core topics in the syllabus, provide:
            1. **ELI5 (Explain Like I'm 5):** A dead-simple, 3-line explanation. No jargon.
            2. **Real-Life Example:** A simple, fun example.
            3. **YouTube Link:** Use this exact format: [📺 Watch Tutorial](https://www.youtube.com/results?search_query=[Replace+With+Topic+Name+With+Plus+Signs])

            Syllabus: {extracted_text}
            """

        elif mode == "test":
            prompt = f"""
            You are an expert examiner. Create a quick, high-yield Self-Test based on the syllabus.
            Include:
            ## 📝 Self-Test Mock Exam
            (Generate exactly 15 MCQs and 8 Short Answer Questions). Leave blank lines between questions.
            ---
            ## 🔐 Answer Key & Weak Point Analyzer
            (Provide correct answers and what topic to revise if they get it wrong).

            Syllabus: {extracted_text}
            """

        response = client.models.generate_content(model='gemini-2.5-flash', contents=prompt)
        return {"content": response.text}

    except Exception as e:
        return {"error": str(e)}